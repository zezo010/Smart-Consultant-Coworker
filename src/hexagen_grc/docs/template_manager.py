"""Template manager for handling DOCX templates with variable mapping."""

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from docx import Document

from ..common.config import settings
from ..common.logging import get_logger
from ..common.schemas import ClientProfile, Language

logger = get_logger(__name__)


class TemplateManager:
    """
    Manages DOCX templates and variable mappings.

    Features:
    - Load mapping from Excel/CSV
    - Validate required variables
    - Map data to placeholders
    - Fill DOCX templates
    """

    def __init__(self, mapping_file: Optional[Path] = None):
        """
        Initialize template manager.

        Args:
            mapping_file: Path to mapping CSV/Excel file
        """
        self.logger = logger
        self.mapping_file = mapping_file or self._get_default_mapping_file()
        self.mapping_data = self._load_mapping()

    def _get_default_mapping_file(self) -> Path:
        """Get default mapping file path."""
        return settings.templates_path / "HexaGen_Policy_Template_Mapping.csv"

    def _load_mapping(self) -> pd.DataFrame:
        """
        Load mapping data from CSV/Excel.

        Returns:
            DataFrame with mapping data
        """
        if not self.mapping_file.exists():
            logger.warning(f"Mapping file not found: {self.mapping_file}")
            return pd.DataFrame()

        try:
            # Try CSV first
            if self.mapping_file.suffix == '.csv':
                df = pd.read_csv(self.mapping_file)
            else:
                df = pd.read_excel(self.mapping_file, sheet_name='Variables')

            logger.info(f"Loaded {len(df)} variable mappings")
            return df

        except Exception as e:
            logger.error(f"Error loading mapping file: {e}")
            return pd.DataFrame()

    def get_placeholder_map(self) -> Dict[str, str]:
        """
        Get dictionary of variable keys to placeholders.

        Returns:
            Dict mapping VariableKey to Placeholder
        """
        if self.mapping_data.empty:
            return {}

        return dict(zip(
            self.mapping_data['VariableKey'],
            self.mapping_data['Placeholder']
        ))

    def get_required_variables(self) -> List[str]:
        """
        Get list of required variable keys.

        Returns:
            List of required VariableKey values
        """
        if self.mapping_data.empty:
            return []

        required = self.mapping_data[self.mapping_data['Required'] == True]
        return required['VariableKey'].tolist()

    def build_variable_data(
        self,
        client_profile: ClientProfile,
        document_info: Dict[str, Any],
        generated_content: Dict[str, Any],
        language: Language = Language.ARABIC
    ) -> Dict[str, Any]:
        """
        Build complete variable data from various sources.

        Args:
            client_profile: Client profile
            document_info: Document metadata
            generated_content: Generated content sections
            language: Target language

        Returns:
            Dictionary of variable values
        """
        # Initialize with default values
        variables = {}

        # Client information
        variables['CLIENT_NAME_AR'] = client_profile.client_name_ar or client_profile.client_name
        variables['CLIENT_NAME_EN'] = client_profile.client_name_en or client_profile.client_name

        # Document information
        variables['POLICY_NAME_AR'] = document_info.get('policy_name_ar', '')
        variables['POLICY_NAME_EN'] = document_info.get('policy_name_en', '')
        variables['DOC_ID'] = document_info.get('doc_id', 'DOC-XXXX-001')
        variables['VERSION'] = document_info.get('version', '1.0')
        variables['CLASSIFICATION'] = document_info.get('classification', 'Internal')
        variables['SHARING'] = document_info.get('sharing_label', 'Orange')
        variables['APPLICABILITY'] = document_info.get('applicability', 'All Departments')
        variables['DOC_TITLE'] = document_info.get('title', '')

        # Dates
        today = datetime.now()
        variables['ISSUE_YEAR'] = str(document_info.get('issue_year', today.year))
        variables['ISSUE_DATE'] = self._format_date(
            document_info.get('issue_date', today),
            language
        )
        variables['NEXT_REVIEW_DATE'] = self._format_date(
            document_info.get('next_review_date'),
            language
        )
        variables['EFFECTIVITY_DATE'] = self._format_date(
            document_info.get('effectivity_date'),
            language
        )

        # Department
        variables['DEPARTMENT_OWNER'] = document_info.get('department_owner', '')

        # Generated content
        variables['POLICY_CONTENT'] = generated_content.get('policy_blocks', '')
        variables['REFERENCES'] = generated_content.get('references', '')
        variables['CONTACT_INFO'] = generated_content.get('contact', '')
        variables['APPENDICES'] = generated_content.get('appendices', '')

        # Approval info
        approvals = document_info.get('approvals', {})
        variables['APPROVAL_NAME'] = approvals.get('name', '')
        variables['APPROVAL_TITLE'] = approvals.get('title', '')
        variables['APPROVAL_DATE'] = self._format_date(
            approvals.get('date'),
            language
        )

        # Metadata
        variables['PREPARED_BY'] = document_info.get('prepared_by', 'HexaGen GRC Team')
        variables['REVIEWED_BY'] = document_info.get('reviewed_by', '')

        return variables

    def _format_date(
        self,
        date: Optional[datetime],
        language: Language = Language.ARABIC
    ) -> str:
        """
        Format date according to language.

        Args:
            date: Date to format
            language: Target language

        Returns:
            Formatted date string
        """
        if not date:
            return ""

        if isinstance(date, str):
            return date

        if language == Language.ARABIC:
            return date.strftime("%d/%m/%Y")
        else:
            return date.strftime("%Y-%m-%d")

    def validate_variables(
        self,
        variables: Dict[str, Any],
        raise_on_missing: bool = False
    ) -> Dict[str, List[str]]:
        """
        Validate that all required variables are present.

        Args:
            variables: Variable data
            raise_on_missing: Whether to raise exception on missing vars

        Returns:
            Dict with 'missing' and 'present' lists
        """
        required = self.get_required_variables()
        present = []
        missing = []

        for var_key in required:
            if var_key in variables and variables[var_key]:
                present.append(var_key)
            else:
                missing.append(var_key)

        if missing:
            logger.warning(f"Missing required variables: {missing}")
            if raise_on_missing:
                raise ValueError(f"Missing required variables: {missing}")

        return {
            'missing': missing,
            'present': present,
            'required_count': len(required),
            'present_count': len(present)
        }

    def fill_template(
        self,
        template_path: Path,
        variables: Dict[str, Any],
        output_path: Path,
        language: Language = Language.ARABIC
    ) -> Path:
        """
        Fill DOCX template with variables.

        Args:
            template_path: Path to template file
            variables: Variable data
            output_path: Path to save filled document
            language: Document language

        Returns:
            Path to filled document
        """
        logger.info(f"Filling template: {template_path}")

        # Load template
        doc = Document(str(template_path))

        # Get placeholder map
        placeholder_map = self.get_placeholder_map()

        # Build replacements (placeholder -> value)
        replacements = {}
        for var_key, placeholder in placeholder_map.items():
            value = variables.get(var_key, '')
            if value:
                replacements[placeholder] = str(value)

        # Replace in paragraphs
        for paragraph in doc.paragraphs:
            for placeholder, value in replacements.items():
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, value)

                    # Apply RTL for Arabic
                    if language == Language.ARABIC:
                        paragraph.paragraph_format.right_to_left = True

        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for placeholder, value in replacements.items():
                            if placeholder in paragraph.text:
                                paragraph.text = paragraph.text.replace(placeholder, value)

                                if language == Language.ARABIC:
                                    paragraph.paragraph_format.right_to_left = True

        # Save filled document
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))

        logger.info(f"Filled template saved: {output_path}")

        return output_path

    def get_variable_info(self, variable_key: str) -> Optional[Dict[str, Any]]:
        """
        Get information about a variable.

        Args:
            variable_key: Variable key

        Returns:
            Variable information or None
        """
        if self.mapping_data.empty:
            return None

        row = self.mapping_data[self.mapping_data['VariableKey'] == variable_key]

        if row.empty:
            return None

        return row.iloc[0].to_dict()

    def list_all_variables(self) -> List[Dict[str, Any]]:
        """
        Get list of all variables.

        Returns:
            List of variable info dictionaries
        """
        if self.mapping_data.empty:
            return []

        return self.mapping_data.to_dict('records')
